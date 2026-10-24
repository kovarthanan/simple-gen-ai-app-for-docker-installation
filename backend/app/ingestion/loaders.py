"""Document loaders: PDF, DOCX, TXT/MD/CSV/JSON, HTML and URLs."""
from __future__ import annotations

import io
import ipaddress
import socket
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import PurePath
from urllib.parse import urljoin, urlparse

import httpx

from app.core.exceptions import GuardrailViolation, IngestionError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".html", ".htm"}


@dataclass
class LoadedPage:
    text: str
    page: int | None = None


@dataclass
class LoadedDocument:
    filename: str
    pages: list[LoadedPage]
    content_type: str | None = None
    source: str = "upload"
    metadata: dict = field(default_factory=dict)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)


def file_extension(filename: str) -> str:
    return PurePath(filename).suffix.lower()


# --- HTML -------------------------------------------------------------------------
class _TextExtractor(HTMLParser):
    _SKIP = {"script", "style", "noscript", "svg", "head"}
    _BLOCK = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section", "article"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip_depth:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    parser = _TextExtractor()
    parser.feed(html)
    return "".join(parser.parts)


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


# --- Files ----------------------------------------------------------------------------
def load_bytes(filename: str, data: bytes, content_type: str | None = None) -> LoadedDocument:
    ext = file_extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise IngestionError(f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}")
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pages = [LoadedPage(text=p.extract_text() or "", page=i + 1) for i, p in enumerate(reader.pages)]
        elif ext == ".docx":
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            pages = [LoadedPage(text="\n\n".join(parts))]
        elif ext in {".html", ".htm"}:
            pages = [LoadedPage(text=html_to_text(_decode(data)))]
        else:
            pages = [LoadedPage(text=_decode(data))]
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError(f"Could not parse '{filename}': {exc}") from exc
    return LoadedDocument(filename=filename, pages=pages, content_type=content_type)


# --- URLs (with SSRF protection) ------------------------------------------------------
def _assert_public_host(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise GuardrailViolation("Only http(s) URLs are allowed")
    try:
        infos = socket.getaddrinfo(parsed.hostname, parsed.port or 443)
    except socket.gaierror as exc:
        raise IngestionError(f"Cannot resolve host '{parsed.hostname}'") from exc
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            raise GuardrailViolation(f"URL resolves to a non-public address ({ip}); blocked")


def load_url(url: str, *, allow_private: bool = False, max_bytes: int = 20 * 1024 * 1024,
             timeout: float = 20.0) -> tuple[LoadedDocument, int]:
    """Fetch a URL (following up to 3 redirects, each re-checked) and load it."""
    current = url
    with httpx.Client(timeout=timeout, follow_redirects=False,
                      headers={"User-Agent": "hybrid-rag-ingest/0.1"}) as client:
        for _ in range(4):
            if not allow_private:
                _assert_public_host(current)
            try:
                resp = client.get(current)
            except httpx.HTTPError as exc:
                raise IngestionError(f"Could not fetch URL: {exc}") from exc
            if resp.is_redirect and resp.headers.get("location"):
                current = urljoin(current, resp.headers["location"])
                continue
            break
        else:
            raise IngestionError("Too many redirects")
    if resp.status_code >= 400:
        raise IngestionError(f"URL returned HTTP {resp.status_code}")
    data = resp.content
    if len(data) > max_bytes:
        raise GuardrailViolation(f"URL content exceeds {max_bytes // (1024 * 1024)} MB")

    ctype = resp.headers.get("content-type", "").split(";")[0].strip().lower()
    name = PurePath(urlparse(current).path).name or urlparse(current).hostname or "page"
    if "pdf" in ctype or name.lower().endswith(".pdf"):
        doc = load_bytes(name if name.endswith(".pdf") else f"{name}.pdf", data, ctype)
    elif "html" in ctype or not ctype:
        doc = LoadedDocument(filename=name, pages=[LoadedPage(text=html_to_text(_decode(data)))], content_type=ctype)
    else:
        doc = LoadedDocument(filename=name, pages=[LoadedPage(text=_decode(data))], content_type=ctype)
    doc.source = current
    return doc, len(data)
